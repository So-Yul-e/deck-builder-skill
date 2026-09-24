#!/usr/bin/env python3
"""deck.py 공통 레이아웃 회귀 검사.

사용자 피드백으로 고친 한·두 줄 정의 목록, 긴 한 줄 제목, 리드-이미지 여백,
세로형 이미지 분할, 리드가 있는 순환 플로우, 카드 값 위계, 콘텐츠 기반 표 높이,
statement 선과 글자 획 정렬이 이후 변경에서 다시 깨지지 않는지 PPTX 구조로 확인한다.
"""

from __future__ import annotations

import math
import sys
import tempfile
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches


# deck.py 와 같은 scripts/ 폴더에 둔다.
SKILL_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(SKILL_DIR))

from deck import BOTTOM, CW, M, T_LEAD, T_VALUE, T_VALUE_TEXT, Deck, measure_pt  # noqa: E402


def _text_shape(slide, text: str):
    expected = " ".join(text.split())
    for shape in slide.shapes:
        actual = " ".join(shape.text.split()) if getattr(shape, "has_text_frame", False) else ""
        if actual == expected:
            return shape
    raise AssertionError(f"텍스트 도형을 찾지 못했습니다: {text}")


def _pictures(slide):
    return [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]


def _tables(slide):
    return [shape for shape in slide.shapes if getattr(shape, "has_table", False)]


def _font_size(shape):
    return shape.text_frame.paragraphs[0].runs[0].font.size.pt


def _make_image(path: Path, width: int, height: int) -> None:
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    for y in range(0, height, max(1, height // 12)):
        draw.rectangle((0, y, width, min(height, y + 4)), fill=(61, 82, 213))
    image.save(path)


def main() -> None:
    with tempfile.TemporaryDirectory(prefix="deck-layout-regression-") as tmp:
        tmp_dir = Path(tmp)
        wide = tmp_dir / "wide.png"
        tall = tmp_dir / "tall.png"
        mobile_paths = [tmp_dir / f"mobile-{idx}.png" for idx in range(3)]
        _make_image(wide, 1600, 700)
        _make_image(tall, 600, 3000)
        for path in mobile_paths:
            _make_image(path, 700, 1400)

        deck = Deck(footer="deck layout regression")
        mixed = deck.deflist(
            "한 줄과 두 줄이 섞여도 행 중심은 일정해야 합니다",
            [
                ("가정", "한 줄 설명은 행 전체의 세로 중앙에 둡니다."),
                ("실행", "두 줄 설명은 첫째 줄과 둘째 줄의 간격을 유지하고\n블록 전체를 행 중앙에 둡니다."),
                ("검증", "라벨과 본문의 기준선이 달라도 시각 중심은 맞춥니다."),
                ("성과", "행 높이를 동일하게 유지해 항목별 무게가 달라지지 않게 합니다."),
            ],
            eyebrow="DEFLIST REGRESSION",
        )
        # eyebrow 세로 정렬 검사용. 막대와 글자의 중심이 같아야 한다.
        deck.deflist(
            "eyebrow 막대와 글자는 같은 중심을 씁니다",
            [("검사", "세로 막대와 글자 상자의 중심이 일치해야 합니다.")],
            eyebrow="eyebrow-align",
        )
        long_title_1 = "주문 접수 자동화는 기존 도구의 한계를 확인한 뒤 구조를 바꿨습니다"
        title_slide_1 = deck.deflist(
            long_title_1,
            [("출발", "반복 업무의 병목을 확인했습니다."), ("전환", "도구의 제약에 맞춰 구조를 바꿨습니다.")],
            lead="실제 한 줄 제목을 두 줄로 오판해 리드만 아래로 내려 보내면 안 됩니다.",
        )
        long_title_2 = "Alpha와 Beta 서비스의 AI 기획 범위를 구분했습니다"
        long_eyebrow = "CASE 07 · ALPHA & BETA SERVICE · PLANNING EVIDENCE"
        long_lead = "Alpha: 사업·서비스 기획과 초기 제품 구축 후 운영팀 이관  |  Beta: PO·UX·데이터 구조 기획, AI 구현 협업"
        title_slide_2 = deck.deflist(
            long_title_2,
            [("Alpha", "모델 요구사항과 사업·서비스 기획 범위를 구분했습니다."), ("Beta", "외부 제품 협업의 역할과 결과를 분리했습니다.")],
            eyebrow=long_eyebrow,
            lead=long_lead,
        )
        shot_slide = deck.shots(
            "이미지는 리드 문장과 붙지 않아야 합니다",
            [str(wide)],
            lead="헤더와 이미지 블록 사이에 공통 호흡을 둡니다.",
        )
        tall_slide = deck.shots(
            "긴 화면은 설명을 밀어내지 않고 패널로 나눕니다",
            [str(tall)],
            captions=["상단에서 하단 순서"],
        )
        loop_lead = "보조 설명이 추가돼도 카드와 순환선은 푸터 위에서 함께 재배치됩니다."
        loop_slide = deck.flow(
            "질문은 답의 유효기간에 맞춰 나누어 제공합니다",
            [("처음 한 번", "평소 습관"), ("후보 보기", "오늘의 맥락"),
             ("방문 후", "방문 만족도"), ("다음 추천", "가중치 반영")],
            lead=loop_lead,
            loop=True,
        )
        concept_cards = deck.cards(
            "입력마다 저장 수명과 추천 역할을 다르게 관리합니다",
            [("최초 질문", "프로필", "사용자가 수정할 때까지"),
             ("오늘 질문", "추천 맥락", "해당 추천에만 연결"),
             ("방문 피드백", "학습 신호", "일일 배치로 반영")],
        )
        mixed_cards = deck.cards(
            "유효 피드백이 쌓일수록 개인 가중치를 확대합니다",
            [("피드백 0건", "0.0 / 0.4 / 0.6", "개인 / 오늘 / 전체"),
             ("피드백 1~4건", "점진 전환", "개인 비중 확대"),
             ("피드백 5건+", "0.5 / 0.3 / 0.2", "개인 취향 우선")],
        )
        numeric_cards = deck.cards(
            "행동 신호는 별점보다 약하게 적용합니다",
            [("다시 보기", "-0.35", "약한 음성 신호"),
             ("중도 종료", "-0.25", "약한 음성 신호"),
             ("장소 저장", "+0.10", "약한 양성 신호")],
        )
        phone_slide = deck.phone_detail(
            "처음 한 번은 평소 이용 습관만 확인합니다",
            str(mobile_paths[0]),
            [("기본 인원", "오늘 인원을 묻기 전 보여줄 기본값입니다."),
             ("방문 경험", "처음 방문과 재방문 추천의 초기 기준입니다."),
             ("이동 방식", "이동 가능한 장소와 동선을 거르는 힌트입니다."),
             ("선호 분위기", "건너뛸 수 있고 나중에 바꿀 수 있습니다.")],
            lead="프로필 값은 다음 추천의 기본값이며 사용자가 바꿀 때까지 유지합니다.",
        )
        dense_table = deck.table(
            "여섯 행 표도 빈 높이까지 부풀리지 않습니다",
            ["구분", "수집 시점", "사용 범위"],
            [[f"항목 {idx}", "추천 전", "다음 추천"] for idx in range(1, 7)],
            lead="본문 줄 수로 행 높이를 정하고 표 전체만 가운데에 둡니다.",
            col_ratio=[2, 2, 3],
        )
        cover_slide = deck.cover(
            "흩어진 운영 데이터를 모아,\n매주 쓰이는 대시보드로 만듭니다",
            "요구 정리 · 데이터 연결 · 화면 구현 · 운영 점검",
        )
        statement_slide = deck.statement(
            "좋은 덱은 결정을 먼저 보이게 합니다.",
            "contact@example.com",
            link_text="Website · example.com",
            link_url="https://example.com",
        )
        two_line_statement = deck.statement(
            "첫 줄의 주장과 선을 맞춥니다.\n둘째 줄의 끝까지 선을 맞춥니다.",
            "두 줄 보조 설명입니다.",
        )
        output = tmp_dir / "layout-regression.pptx"
        deck.save(output)

        # deflist: 모든 라벨/본문은 행 중앙, 두 줄 본문만 1.32 행간.
        for text in ("가정", "실행", "검증", "성과"):
            assert _text_shape(mixed, text).text_frame.vertical_anchor == MSO_ANCHOR.MIDDLE
        one_line = _text_shape(mixed, "한 줄 설명은 행 전체의 세로 중앙에 둡니다.")
        two_line = _text_shape(
            mixed,
            "두 줄 설명은 첫째 줄과 둘째 줄의 간격을 유지하고\n블록 전체를 행 중앙에 둡니다.",
        )
        assert one_line.text_frame.vertical_anchor == MSO_ANCHOR.MIDDLE
        assert one_line.text_frame.paragraphs[0].line_spacing is None
        assert two_line.text_frame.vertical_anchor == MSO_ANCHOR.MIDDLE
        assert math.isclose(two_line.text_frame.paragraphs[0].line_spacing, 1.32, rel_tol=0.01)

        # 동일한 행 높이는 행 사이 구분선의 y 간격으로 검증한다.
        separators = sorted(
            shape.top
            for shape in mixed.shapes
            if shape.left == M and shape.width == CW and shape.height == 9525 and shape.top < Inches(6.2)
        )
        assert len(separators) == 3, f"행 구분선 수가 다릅니다: {len(separators)}"
        distances = [b - a for a, b in zip(separators, separators[1:])]
        assert max(distances) - min(distances) <= 2, f"행 높이가 일정하지 않습니다: {distances}"

        # 26pt 로 줄이면 한 줄에 들어가는 긴 제목은 둘 다 한 줄 높이를 유지한다.
        for slide, title in ((title_slide_1, long_title_1), (title_slide_2, long_title_2)):
            title_shape = _text_shape(slide, title)
            assert title_shape.height == Inches(0.8), f"한 줄 제목을 두 줄로 오판했습니다: {title}"

        # 긴 eyebrow는 고정 7in에 가두지 않고 본문 폭을 쓰며, 긴 리드는 실측한
        # 문단으로 나눠 LibreOffice에서도 좌우로 역류하지 않게 한다.
        eyebrow_shape = _text_shape(title_slide_2, long_eyebrow)
        assert eyebrow_shape.width == CW - Inches(0.22)
        lead_shape = _text_shape(title_slide_2, long_lead)
        lead_paragraphs = [p.text for p in lead_shape.text_frame.paragraphs]
        assert len(lead_paragraphs) == 2, f"긴 리드 줄 수가 다릅니다: {lead_paragraphs}"
        lead_width_pt = CW / 12700
        assert all(measure_pt(line, T_LEAD) * 1.12 <= lead_width_pt for line in lead_paragraphs)

        # shots: 한 줄 리드 하단 이후 기본 0.12in + 이미지 전용 0.16in을 확보한다.
        lead = _text_shape(shot_slide, "헤더와 이미지 블록 사이에 공통 호흡을 둡니다.")
        first_picture = _pictures(shot_slide)[0]
        expected_top = lead.top + lead.height + Inches(0.12) + Inches(0.16)
        assert first_picture.top >= expected_top - 2, "리드와 이미지 사이 공통 여백이 사라졌습니다."

        # 5:1 세로 화면은 한 장으로 축소하지 않고 4개 패널로 분할한다.
        assert len(_pictures(tall_slide)) == 4, "긴 단일 이미지 자동 분할 규칙이 깨졌습니다."

        loop_lead_shape = _text_shape(loop_slide, loop_lead)
        first_step = _text_shape(loop_slide, "1. 처음 한 번")
        first_desc = _text_shape(loop_slide, "평소 습관")
        last_step = _text_shape(loop_slide, "4. 다음 추천")
        last_desc = _text_shape(loop_slide, "가중치 반영")
        loop_label = _text_shape(loop_slide, "→  다음 방문에 반영")
        assert first_step.top >= loop_lead_shape.top + loop_lead_shape.height
        assert loop_label.top + loop_label.height <= BOTTOM

        # flow 카드의 제목·설명 묶음은 각 칸 중심을 기준으로 배치한다. 고정된
        # 위/아래 좌표를 쓰면 보조글 줄 수에 따라 묶음 전체가 아래로 쏠린다.
        flow_cards = [
            shape for shape in loop_slide.shapes
            if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
            and shape.width > Inches(2.0) and shape.height > Inches(1.0)
            and shape.top < BOTTOM
        ]
        assert len(flow_cards) == 4, f"flow 카드 수가 다릅니다: {len(flow_cards)}"
        for card, title_shape, desc_shape in zip(
            sorted(flow_cards, key=lambda shape: shape.left),
            (first_step, _text_shape(loop_slide, "2. 후보 보기"),
             _text_shape(loop_slide, "3. 방문 후"), last_step),
            (first_desc, _text_shape(loop_slide, "오늘의 맥락"),
             _text_shape(loop_slide, "방문 만족도"), last_desc),
        ):
            text_top = min(title_shape.top, desc_shape.top)
            text_bottom = max(title_shape.top + title_shape.height,
                              desc_shape.top + desc_shape.height)
            text_center = (text_top + text_bottom) / 2
            card_center = card.top + card.height / 2
            assert abs(text_center - card_center) <= Inches(0.04), (
                f"flow 카드 내부 묶음이 칸 중심에서 벗어났습니다: "
                f"{(text_center - card_center) / 914400:.3f}in"
            )

        for slide, value in ((concept_cards, "프로필"), (mixed_cards, "0.0 / 0.4 / 0.6")):
            assert _font_size(_text_shape(slide, value)) <= T_VALUE_TEXT
        assert _font_size(_text_shape(numeric_cards, "-0.35")) <= T_VALUE

        phone_lead = _text_shape(
            phone_slide, "프로필 값은 다음 추천의 기본값이며 사용자가 바꿀 때까지 유지합니다."
        )
        phone_picture = _pictures(phone_slide)[0]
        assert phone_picture.top >= phone_lead.top + phone_lead.height
        assert phone_picture.top + phone_picture.height <= BOTTOM
        guard_deck = Deck()
        try:
            guard_deck.shots("모바일 화면 세 개", [str(path) for path in mobile_paths])
        except ValueError as exc:
            assert "phone_detail()" in str(exc)
        else:
            raise AssertionError("세로형 모바일 화면 3개가 축소 갤러리로 생성됐습니다.")

        table_shape = _tables(dense_table)[0]
        table = table_shape.table
        body_heights = [table.rows[idx].height for idx in range(1, len(table.rows))]
        assert max(body_heights) <= Inches(0.46) + 2
        assert table_shape.top + table_shape.height <= BOTTOM

        # 두 줄 표지 카피는 전용 높이를 쓰며 서브카피와 겹치지 않는다.
        cover_title = _text_shape(
            cover_slide, "흩어진 운영 데이터를 모아,\n매주 쓰이는 대시보드로 만듭니다"
        )
        cover_subtitle = _text_shape(cover_slide, "요구 정리 · 데이터 연결 · 화면 구현 · 운영 점검")
        assert cover_title.top + cover_title.height <= cover_subtitle.top

        # 마지막 CTA는 주소를 보이는 텍스트와 실제 하이퍼링크가 함께 있어야 한다.
        statement_link = _text_shape(statement_slide, "Website · example.com")
        statement_contact = _text_shape(statement_slide, "contact@example.com")
        assert statement_link.top + statement_link.height <= statement_contact.top
        hyperlink_runs = [
            run for paragraph in statement_link.text_frame.paragraphs for run in paragraph.runs
            if run.hyperlink.address
        ]
        assert [run.hyperlink.address for run in hyperlink_runs] == ["https://example.com"]

        # statement 선은 텍스트박스가 아니라 Pretendard의 실제 glyph 영역에 맞춘다.
        # 박스 전체 높이를 쓰면 한 줄은 위로 0.18in, 두 줄은 0.12in 튀어나온다.
        for slide, title, top_inset, bottom_outset in (
            (statement_slide, "좋은 덱은 결정을 먼저 보이게 합니다.", 0.18, 0.03),
            (two_line_statement, "첫 줄의 주장과 선을 맞춥니다.\n둘째 줄의 끝까지 선을 맞춥니다.", 0.12, 0.08),
        ):
            title_shape = _text_shape(slide, title)
            bars = [
                shape for shape in slide.shapes
                if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
                and shape.left == M and shape.width == Inches(0.06)
                and shape.top < BOTTOM
            ]
            assert len(bars) == 1, f"statement 세로 막대 수가 다릅니다: {len(bars)}"
            bar = bars[0]
            assert abs(bar.top - (title_shape.top + Inches(top_inset))) <= 2
            assert abs(
                (bar.top + bar.height)
                - (title_shape.top + title_shape.height + Inches(bottom_outset))
            ) <= 2

        # 3줄 설명은 조용히 밀어 넣지 않고 생성 단계에서 차단한다.
        overflow_deck = Deck()
        try:
            overflow_deck.deflist("overflow", [("항목", "첫째 줄\n둘째 줄\n셋째 줄")])
        except ValueError as exc:
            assert "최대 2줄" in str(exc)
        else:
            raise AssertionError("deflist 3줄 설명이 차단되지 않았습니다.")

        # eyebrow: 세로 막대와 글자가 같은 세로 중심을 공유해야 한다.
        # 한 줄 텍스트에 기본 행간(1.35)이 붙으면 늘어난 leading이 글자 위에 쌓여
        # MIDDLE 앵커가 줄 상자를 가운데 둬도 글자만 아래로 쏠린다.
        # 실측 2026-08-17: 3.1pt 어긋나 있었다.
        eyebrow_slide = None
        for slide in Presentation(output).slides:
            hit = [sh for sh in slide.shapes
                   if getattr(sh, "has_text_frame", False)
                   and " ".join(sh.text.split()).strip() == "EYEBROW-ALIGN"]
            if hit:
                eyebrow_slide, eyebrow_box = slide, hit[0]
                break
        assert eyebrow_slide is not None, "eyebrow 검사용 슬라이드를 찾지 못했습니다."
        bars = [sh for sh in eyebrow_slide.shapes
                if not getattr(sh, "has_text_frame", False) or not sh.text.strip()]
        bars = [b for b in bars if b.width < Inches(0.12) and b.height < Inches(0.6)]
        assert bars, "eyebrow 세로 막대를 찾지 못했습니다."
        bar = bars[0]
        assert abs((bar.top + bar.height / 2) - (eyebrow_box.top + eyebrow_box.height / 2)) <= Inches(0.005), \
            "eyebrow 막대와 글자 상자의 세로 중심이 어긋났습니다."
        for paragraph in eyebrow_box.text_frame.paragraphs:
            assert paragraph.line_spacing is None, \
                "eyebrow 한 줄 텍스트에 행간이 붙으면 글자가 아래로 쏠립니다."

        assert output.exists() and output.stat().st_size > 0
        print("PASS: deck layout regression (deflist/title/flow/cards/table/shots/statement/eyebrow/overflow)")


if __name__ == "__main__":
    main()
