#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Add subject imagery to an inherited deck without rebuilding its content."""
import argparse
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

WHITE=RGBColor(255,255,255)


def picture_cover(slide,path,x,y,w,h,focus_x=.5):
    """Native PPT picture crop; preserve original raster asset and aspect ratio."""
    with Image.open(path) as image:
        iw,ih=image.size
    target=w/h
    source=iw/ih
    picture=slide.shapes.add_picture(str(path), Inches(x), Inches(y),width=Inches(w),height=Inches(h))
    if source>target:
        cropped=1-target/source
        picture.crop_left=cropped*focus_x
        picture.crop_right=cropped*(1-focus_x)
    else:
        cropped=1-source/target
        picture.crop_top=picture.crop_bottom=cropped/2
    tree=slide.shapes._spTree
    tree.remove(picture._element)
    tree.insert(2,picture._element)
    return picture


def place(shape,x,y,w,h,size=None,color=None):
    shape.left,shape.top,shape.width,shape.height=map(Inches,(x,y,w,h))
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            for run in p.runs:
                if size is not None: run.font.size=Pt(size)
                if color is not None: run.font.color.rgb=color


def build(baseline,image,output):
    if baseline.resolve()==output.resolve():
        raise ValueError('Build a separate candidate; do not overwrite the baseline')
    prs=Presentation(baseline)
    if len(prs.slides)!=6:
        raise ValueError('This demonstration expects the inherited six-slide art-direction deck')
    cover=prs.slides[0]
    original=list(cover.shapes)
    picture_cover(cover,image,0,0,13.333,7.5)
    for s in original:
        if not s.has_text_frame: continue
        if '같은 틀에' in s.text:
            s.text_frame.text='같은 틀에 넣으면\n다른 내용도 같아 보인다.'
            s.text_frame.paragraphs[1].text='다른 내용도 \v같아 보인다.'
            for paragraph in s.text_frame.paragraphs:
                paragraph.line_spacing=1.15
                for run in paragraph.runs:
                    run.font.name='Pretendard'
                    run.font.bold=True
            place(s,.75,1.45,5.0,3.4,36,WHITE)
        elif '정렬은' in s.text:
            place(s,.75,5.15,5,1.1,21,WHITE)
        else:
            for p in s.text_frame.paragraphs:
                for run in p.runs: run.font.color.rgb=WHITE
    for s in original:
        if not s.has_text_frame or not s.text:
            if s.width>Inches(6):
                s.width=Inches(4.8)
    # Keep the full-screen image only on the opening; the body uses a material detail.
    body=prs.slides[1]
    original=list(body.shapes)
    picture_cover(body,image,8.5,0,4.833,7.5,focus_x=1.0)
    for s in original:
        if not s.has_text_frame or not s.text:
            if s.height>Inches(1): place(s,3.65,2.8,.035,2.65)
            continue
        value=s.text
        if '색을 바꿔도' in value: place(s,.75,.65,7.05,1.65,30)
        elif value=='입력': place(s,.75,3.0,2,.4,13)
        elif value=='표현': place(s,4.05,3.0,3,.4,13)
        elif value=='절차 / 비교 / 실물': place(s,.75,3.6,2.65,1.7,21)
        elif value=='흐름 / 대조 / 화면': place(s,4.05,3.6,3.7,1.7,21)
        elif '절차는 흐름으로' in value: place(s,.75,5.8,7.1,1.0,17)
        elif value=='2': place(s,7.5,7.02,.5,.25,10)
        else: place(s,.75,7.03,6.7,.23,10)
    for n in (0,1):
        notes=prs.slides[n].notes_slide.notes_text_frame
        notes.text=notes.text+'\nImage: generated illustrative editorial workbench; not evidence of an actual workspace. Asset: '+image.name
    output.parent.mkdir(parents=True,exist_ok=True)
    prs.save(output)
    print(output)


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--baseline',type=Path,required=True)
    parser.add_argument('--image',type=Path,required=True)
    parser.add_argument('--output',type=Path,required=True)
    args=parser.parse_args()
    build(args.baseline,args.image,args.output)
