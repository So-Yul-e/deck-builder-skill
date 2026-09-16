from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from PIL import Image
from pptx.enum.shapes import MSO_SHAPE_TYPE


REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "deck" / "scripts"))

from deck import Deck, PALETTES, SW  # noqa: E402


CONCEPTS = ("report", "poster", "editorial", "showcase")


def _texts(slide) -> list[str]:
    return [shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False)]


def _text_box(slide, text: str):
    return next(shape for shape in slide.shapes if getattr(shape, "text", "") == text)


def _slide_signature(slide) -> list[tuple[int, int, int, int, int, str]]:
    """Compare the editable layout, avoiding volatile PPT relationship IDs."""
    return [
        (shape.shape_type, shape.left, shape.top, shape.width, shape.height,
         shape.text if getattr(shape, "has_text_frame", False) else "")
        for shape in slide.shapes
    ]


class ConceptDeckTests(unittest.TestCase):
    def test_all_concepts_build_cover_cards_and_chart_without_losing_content(self) -> None:
        for concept in CONCEPTS:
            with self.subTest(concept=concept):
                deck = Deck(concept=concept, footer="제품팀")
                deck.cover("한 문장 결론", "보조 설명", "2026.09")
                deck.cards("핵심 지표", [("전환", "42%", "전주 대비 +7%")])
                deck.chart("월별 성장", [("7월", 10), ("8월", 18)])

                self.assertEqual(3, len(deck.prs.slides))
                self.assertTrue({"한 문장 결론", "보조 설명", "2026.09"}.issubset(_texts(deck.prs.slides[0])))
                self.assertTrue({"핵심 지표", "전환", "42%", "전주 대비 +7%"}.issubset(_texts(deck.prs.slides[1])))
                self.assertIn("월별 성장", _texts(deck.prs.slides[2]))
                chart = next(shape.chart for shape in deck.prs.slides[2].shapes
                             if shape.shape_type == MSO_SHAPE_TYPE.CHART)
                self.assertEqual(["7월", "8월"], [point.label for point in chart.plots[0].categories])

    def test_concepts_produce_distinct_cover_layouts(self) -> None:
        title_positions = {}
        for concept in CONCEPTS:
            deck = Deck(concept=concept)
            slide = deck.cover("동일한 제목", "같은 부제", "같은 메타")
            title = _text_box(slide, "동일한 제목")
            title_positions[concept] = (title.left, title.top, title.width, title.height)

        self.assertEqual(
            len(CONCEPTS), len(set(title_positions.values())),
            f"컨셉별 표지 제목 레이아웃이 실제로 달라야 한다: {title_positions}",
        )

    def test_showcase_cover_places_image_on_right_and_preserves_aspect_ratio(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            image_path = Path(directory) / "wide.png"
            Image.new("RGB", (400, 200), "white").save(image_path)

            deck = Deck(concept="showcase")
            slide = deck.cover("제품 소개", image=image_path)
            picture = next(shape for shape in slide.shapes
                           if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)

        self.assertGreaterEqual(picture.left, SW // 2)
        self.assertAlmostEqual(2.0, picture.width / picture.height, places=4)

    def test_each_concept_accepts_a_cover_image(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            image_path = Path(directory) / "cover.png"
            Image.new("RGB", (300, 200), "white").save(image_path)
            for concept in CONCEPTS:
                with self.subTest(concept=concept):
                    slide = Deck(concept=concept).cover("이미지 표지", image=image_path)
                    self.assertTrue(any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE
                                        for shape in slide.shapes))

    def test_explicit_palette_has_priority_over_concept_preset(self) -> None:
        deck = Deck(palette="navy", concept="poster")
        self.assertEqual(PALETTES["navy"]["deep"], str(deck.C["deep"]))

    def test_unknown_concept_is_rejected(self) -> None:
        with self.assertRaisesRegex(ValueError, "concept"):
            Deck(concept="neon-space-opera")

    def test_default_api_matches_explicit_report_and_instances_do_not_leak_state(self) -> None:
        legacy = Deck()
        explicit = Deck(concept="report")
        for deck in (legacy, explicit):
            deck.cover("제목", "부제", "메타")
            deck.cards("지표", [("매출", "10억", "목표 초과")])
            deck.chart("성장", [("1분기", 1), ("2분기", 2)])

        self.assertEqual(
            [_slide_signature(slide) for slide in legacy.prs.slides],
            [_slide_signature(slide) for slide in explicit.prs.slides],
        )

        poster = Deck(concept="poster")
        poster.cover("포스터")
        fresh_report = Deck()
        fresh_slide = fresh_report.cover("보고서")
        expected_report = Deck(concept="report").cover("보고서")
        self.assertEqual(_slide_signature(expected_report), _slide_signature(fresh_slide))


if __name__ == "__main__":
    unittest.main()
