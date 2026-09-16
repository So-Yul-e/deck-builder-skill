from pathlib import Path
import sys
import tempfile
import unittest
from PIL import Image
from pptx.enum.shapes import MSO_SHAPE_TYPE

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'deck'/'scripts'))
from deck import Deck


class ComposeTests(unittest.TestCase):
    def test_text_remains_editable_and_complete_with_named_colors(self):
        d=Deck()
        s=d.compose([dict(kind='text',text='내용의 관계를 먼저 봅니다.',x=.8,y=1,w=4,h=2,size=24,color='key')])
        shape=s.shapes[0]
        self.assertTrue(shape.has_text_frame)
        self.assertEqual('내용의 관계를 먼저 봅니다.',shape.text.replace('\n',''))
        self.assertEqual(str(d.C['key']),str(shape.text_frame.paragraphs[0].runs[0].font.color.rgb))

    def test_invalid_element_leaves_no_partial_slide(self):
        valid=dict(kind='rect',x=1,y=1,w=2,h=1,color='key')
        for invalid in (dict(valid,x=13),dict(valid,color='#ffffff'),dict(valid,w=float('nan')),
                        dict(kind='text',text='넘치는 긴 내용',x=1,y=1,w=1,h=.1,size=50)):
            d=Deck()
            with self.assertRaises(ValueError):
                d.compose([valid,invalid])
            self.assertEqual(0,len(d.prs.slides))

    def test_native_rectangles_override_inherited_theme_effects(self):
        s=Deck().compose([dict(kind='rect',x=1,y=1,w=2,h=1,color='key')])
        effects=s.shapes[0]._element.spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst')
        self.assertIsNotNone(effects)
        self.assertEqual(0,len(effects))
        self.assertTrue(all(r.get('idx') == '0' for r in s.shapes[0]._element.xpath('.//a:effectRef')))

    def test_image_fits_box_without_distortion(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'screen.png'
            Image.new('RGB',(400,200),'white').save(path)
            s=Deck().compose([dict(kind='image',path=path,x=1,y=1,w=3,h=3)])
            picture=s.shapes[0]
            self.assertEqual(MSO_SHAPE_TYPE.PICTURE,picture.shape_type)
            self.assertAlmostEqual(2,picture.width/picture.height)
            self.assertAlmostEqual(1.75,picture.top/914400)


if __name__=='__main__':
    unittest.main()
